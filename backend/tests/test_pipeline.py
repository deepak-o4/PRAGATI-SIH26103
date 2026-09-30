import io, unittest, tempfile
from datetime import date
from pathlib import Path
from tests.helpers import *
from app.ingestion.loader import *
from app.analytics.portfolio import dashboard, bottlenecks, data_quality, enrich
from app.copilot.engine import answer
from app.copilot.documents import extract_text, extract_entities, TfidfIndex
from app.reports.generate import executive_csv, executive_xlsx, executive_pdf, project_pdf
from app.ml.delay_model import features_for, predict, train, build_training_set, FEATURES, FALLBACK_VERSION

SAMPLE = Path(__file__).resolve().parents[2] / "data" / "sample"


def load_demo():
    projs, r1 = ingest_projects((SAMPLE / "projects.csv").read_text(encoding="utf-8"))
    codes = {p.project_code for p in projs}
    sn, r2 = ingest_snapshots((SAMPLE / "project_snapshots.csv").read_text(encoding="utf-8"), codes)
    ms, r3 = ingest_milestones((SAMPLE / "project_milestones.csv").read_text(encoding="utf-8"), codes)
    iss, r4 = ingest_issues((SAMPLE / "project_issues.csv").read_text(encoding="utf-8"), codes)
    return assemble(projs, sn, ms, iss), (r1, r2, r3, r4)


class TestIngestion(unittest.TestCase):
    def test_demo_loads_clean_and_labelled(self):
        ps, reps = load_demo()
        self.assertEqual(len(ps), 60)
        for r in reps: self.assertEqual(r.rejected, [], r.kind)
        self.assertTrue(all(p.origin == DataOrigin.SYNTHETIC_DEMO for p in ps))
        self.assertTrue(all(p.project_code.startswith("DEMO-") for p in ps))

    def test_rejections(self):
        csv_text = ("project_code,project_name,original_cost_crore,physical_progress_pct,start_date,original_end_date\n"
                    "A1,ok,100,50,2020-01-01,2025-01-01\n"
                    "A2,neg cost,-5,10,2020-01-01,2025-01-01\n"
                    "A3,bad pct,100,120,2020-01-01,2025-01-01\n"
                    "A4,bad dates,100,10,2025-01-01,2020-01-01\n"
                    "A1,dup,100,10,2020-01-01,2025-01-01\n"
                    ",no code,100,10,2020-01-01,2025-01-01\n"
                    "A7,junk num,abc,10,2020-01-01,2025-01-01\n")
        ps, rep = ingest_projects(csv_text)
        self.assertEqual([p.project_code for p in ps], ["A1"])
        self.assertEqual(len(rep.rejected), 6)
        joined = " ".join(e for r in rep.rejected for e in r["errors"])
        for needle in ("negative", "out of range", "precedes", "duplicate", "missing project_code", "not a number"):
            self.assertIn(needle, joined)

    def test_missing_optional_stays_null(self):
        ps, _ = ingest_projects("project_code,project_name\nX1,Only name\n")
        self.assertIsNone(ps[0].revised_cost_crore); self.assertIsNone(ps[0].physical_progress_pct)

    def test_missing_required_columns(self):
        ps, rep = ingest_projects("foo,bar\n1,2\n")
        self.assertEqual(ps, []); self.assertTrue(rep.rejected)

    def test_paimana_style_aliases_and_formats(self):
        csv_text = "Project Code,Project Name,Sector,Line Ministry,Original Cost (Rs. Crore),Revised Cost,Expenditure,Anticipated Completion Date,Physical Progress (%)\n" \
                   "705728,Big Project,Railways,Ministry of Railways,\"1,08,000\",\"1,20,000.5\",90396,31/03/2028,59.1\n"
        # Indian digit grouping commas are stripped; the value 108000
        ps, rep = ingest_projects(csv_text)
        self.assertEqual(rep.rejected, [])
        p = ps[0]
        self.assertEqual(p.original_cost_crore, 108000.0); self.assertEqual(p.revised_end_date, date(2028, 3, 31))
        self.assertEqual(p.physical_progress_pct, 59.1)

    def test_snapshot_validation_and_dedup(self):
        t = ("project_code,snapshot_month,physical_progress_pct\nA,2026-01,10\nA,2026-01,12\nA,2026-02,11\nB,2026-01,5\nA,2026-03,150\n")
        by, rep = ingest_snapshots(t, {"A"})
        self.assertEqual(len(by["A"]), 2)
        self.assertEqual(len(rep.rejected), 2)
        self.assertTrue(any("duplicate" in w for w in rep.warnings))
        self.assertTrue(any("decreased" in w for w in rep.warnings))


class TestPortfolio(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ps, _ = load_demo()

    def test_dashboard_consistency(self):
        d = dashboard(self.ps, AS_OF)
        self.assertEqual(d["kpis"]["total_projects"], 60)
        self.assertTrue(d["demo_data_present"])
        self.assertEqual(sum(x["count"] for x in d["risk_distribution"]), 60)
        self.assertEqual(sum(x["projects"] for x in d["by_sector"]), 60)
        self.assertAlmostEqual(d["kpis"]["total_original_cost_crore"], round(sum(p.original_cost_crore for p in self.ps), 2), places=1)
        for dist in ("progress_distribution", "cost_overrun_distribution", "schedule_delay_distribution"):
            self.assertEqual(sum(x["count"] for x in d[dist]), 60, dist)

    def test_risk_spread_is_meaningful(self):
        bands = {x["bucket"]: x["count"] for x in dashboard(self.ps, AS_OF)["risk_distribution"]}
        self.assertGreaterEqual(sum(1 for v in bands.values() if v > 0), 3, bands)

    def test_bottlenecks_from_data_only(self):
        b = bottlenecks(self.ps, AS_OF)
        open_total = sum(1 for p in self.ps for i in p.issues if i.is_open)
        self.assertEqual(sum(x["open_issues"] for x in b), open_total)

    def test_data_quality(self):
        dq = data_quality(self.ps, AS_OF)
        self.assertEqual(dq["total_records"], 60); self.assertEqual(dq["duplicate_project_codes"], 0)
        bad = self.ps + [mk("DEMO-0001", physical_progress_pct=150.0, original_cost_crore=-1.0)]
        dq2 = data_quality(bad, AS_OF)
        self.assertEqual(dq2["duplicate_project_codes"], 1); self.assertEqual(dq2["invalid_progress_values"], 1); self.assertEqual(dq2["invalid_costs"], 1)


class TestCopilot(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ps, _ = load_demo()

    def test_intents(self):
        for q, intent in [("Which projects are at critical risk?", "critical_projects"),
                          ("Which projects have the highest cost overruns?", "cost_overrun"),
                          ("Which projects have stalled progress?", "stalled"),
                          ("Show delayed projects in Uttar Pradesh", "delayed"),
                          ("Which projects have high expenditure but low physical progress?", "cost_progress_imbalance"),
                          ("Which sectors have the highest average risk?", "sector_risk"),
                          ("What projects are approaching their revised completion date?", "approaching_deadline"),
                          ("Why is DEMO-0005 at risk?", "explain_risk"),
                          ("Compare DEMO-0001 and DEMO-0002", "compare"),
                          ("Summarize DEMO-0003 for an executive review", "project_summary")]:
            r = answer(q, self.ps, AS_OF)
            self.assertEqual(r["intent"], intent, q)
            self.assertTrue(r["answer"])
            self.assertIsNotNone(r["demo_data_notice"])

    def test_numbers_come_from_data(self):
        r = answer("Which projects are at critical risk?", self.ps, AS_OF)
        n = sum(1 for p in self.ps if enrich(p, AS_OF)["risk_band"] == "CRITICAL" and p.status.value not in ("COMPLETED", "CANCELLED"))
        self.assertIn(f"{n} active project", r["answer"])

    def test_delayed_state_filter(self):
        r = answer("Show delayed projects in Uttar Pradesh", self.ps, AS_OF)
        self.assertTrue(all(row["state"] == "Uttar Pradesh" for row in r["table"]))

    def test_unsupported_does_not_hallucinate(self):
        r = answer("What is the weather like?", self.ps, AS_OF)
        self.assertEqual(r["intent"], "unsupported"); self.assertEqual(r["table"], [])

    def test_empty(self):
        self.assertEqual(answer("critical projects", [], AS_OF)["intent"], "empty")

    def test_narrator_only_rephrases(self):
        r = answer("how many projects", self.ps, AS_OF, narrator=lambda t: t.upper())
        self.assertEqual(r["narrated"], r["answer"].upper())

    def test_retrieval_fallback(self):
        idx = TfidfIndex(); idx.add("review.txt", "The land acquisition for the bypass section is delayed pending court stay order.")
        r = answer("tell me about the court stay order", self.ps, AS_OF, retriever=idx.search)
        self.assertEqual(r["intent"], "document_retrieval"); self.assertTrue(any(s["type"] == "document" for s in r["sources"]))


class TestDocuments(unittest.TestCase):
    def test_txt_and_entities(self):
        txt = "Project DEMO-0001 revised to ₹1,200 crore. Progress 59.1%. Target 2027-03-31."
        e = extract_entities(extract_text("a.txt", txt.encode()), {"DEMO-0001", "DEMO-9999"})
        self.assertEqual(e["project_references"], ["DEMO-0001"]); self.assertIn("59.1%", e["percentages"]); self.assertIn("2027-03-31", e["dates"])

    def test_rejects_bad_ext(self):
        with self.assertRaises(ValueError): extract_text("x.exe", b"MZ")

    def test_docx_xlsx_roundtrip(self):
        import docx, openpyxl
        d = docx.Document(); d.add_paragraph("Milestone slipped"); b = io.BytesIO(); d.save(b)
        self.assertIn("Milestone slipped", extract_text("r.docx", b.getvalue()))
        wb = openpyxl.Workbook(); wb.active.append(["a", 1]); b2 = io.BytesIO(); wb.save(b2)
        self.assertIn("a | 1", extract_text("r.xlsx", b2.getvalue()))


class TestReports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ps, _ = load_demo()

    def test_pdf_csv_xlsx(self):
        pdf = executive_pdf(self.ps, AS_OF); self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertTrue(project_pdf(self.ps[0], AS_OF).startswith(b"%PDF"))
        csv_b = executive_csv(self.ps, AS_OF).decode(); self.assertEqual(len(csv_b.strip().splitlines()), 61)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(executive_xlsx(self.ps, AS_OF)))
        self.assertEqual(wb.sheetnames, ["Summary", "Projects", "Bottlenecks"]); self.assertEqual(wb["Projects"].max_row, 61)
        self.assertIn("SYNTHETIC_DEMO", wb["Summary"]["A2"].value)


class TestML(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ps, _ = load_demo()

    def test_features(self):
        f = features_for(self.ps[0], AS_OF)
        self.assertEqual(set(f), set(FEATURES))
        import math
        self.assertTrue(math.isnan(features_for(mk(revised_cost_crore=None, original_cost_crore=None), AS_OF)["cost_overrun_pct"]))

    def test_fallback_prediction_labelled(self):
        r = predict(self.ps[0], AS_OF)
        self.assertEqual(r["model_version"], FALLBACK_VERSION); self.assertFalse(r["calibrated"])
        self.assertTrue(0 <= r["delay_probability"] <= 1)

    def test_no_training_without_enough_labels(self):
        with tempfile.TemporaryDirectory() as d:
            rep = train(self.ps, d)
            self.assertFalse(rep.trained); self.assertIn("insufficient", rep.reason)
            self.assertFalse((Path(d) / "delay_model.joblib").exists())

    def test_training_mechanics_on_synthetic_labels_only_when_forced(self):
        """Exercise train/load/predict/explain plumbing with a tiny hand-built dataset. NOT evidence of model quality."""
        import numpy as np, joblib
        from sklearn.ensemble import HistGradientBoostingClassifier
        rng = np.random.default_rng(0)
        A = rng.normal(size=(200, len(FEATURES))); y = (A[:, 2] + 0.3 * rng.normal(size=200) > 0).astype(int)
        clf = HistGradientBoostingClassifier(max_iter=30, random_state=0).fit(A, y)
        model = {"model": clf, "features": FEATURES, "train_means": A.mean(0).tolist(), "version": "test-plumbing"}
        with tempfile.TemporaryDirectory() as d:
            joblib.dump(model, Path(d) / "delay_model.joblib")
            from app.ml.delay_model import load_model
            m = load_model(d); self.assertEqual(m["version"], "test-plumbing")
            r = predict(self.ps[0], AS_OF, m)
            self.assertEqual(r["model_version"], "test-plumbing"); self.assertTrue(r["top_factors"])
            self.assertTrue(all(k in r["top_factors"][0] for k in ("feature", "impact", "direction")))

    def test_training_set_labels_use_later_real_observations(self):
        p = mk(snapshots=tuple(Snapshot(date(2025, m, 1), physical_progress_pct=float(m * 5), revised_end_date=date(2027, 1, 1) if m < 8 else date(2027, 9, 1)) for m in range(1, 13)))
        X, y, meta = build_training_set([p], horizon_days=60)
        lab = {m[1]: v for m, v in zip(meta, y)}
        self.assertEqual(lab["2025-03-01"], 1)               # end date was later pushed out ~8 months
        self.assertEqual(lab["2025-09-01"], 0)               # already reflected the later date -> no new slip
        self.assertNotIn(("T-1", "2025-12-01"), meta)        # last snapshot has no future -> excluded

if __name__ == "__main__":
    unittest.main()
