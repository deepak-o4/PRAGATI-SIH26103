"""Document intelligence: text extraction, light entity extraction, TF-IDF retrieval (no external service)."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

ALLOWED_EXT = {".pdf", ".docx", ".xlsx", ".csv", ".txt"}
MAX_BYTES = 25 * 1024 * 1024


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ValueError(f"unsupported file type {ext!r}; allowed: {sorted(ALLOWED_EXT)}")
    if len(data) > MAX_BYTES:
        raise ValueError("file too large")
    if ext == ".txt":
        return data.decode("utf-8", errors="replace")
    if ext == ".csv":
        return "\n".join(", ".join(r) for r in csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(data)).pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return "\n".join(parts)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## {ws.title}")
            for row in ws.iter_rows(values_only=True):
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)
    raise AssertionError


_CODE = re.compile(r"\b[A-Z]{2,6}-\d{2,8}\b|\b\d{5,8}\b")
_DATE = re.compile(r"\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}[/-]\d{1,2}[/-]\d{4}\b|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[ -]\d{4}\b")
_COST = re.compile(r"(?:₹|Rs\.?\s*)\s*[\d,]+(?:\.\d+)?\s*(?:crore|cr\.?)?", re.I)
_PCT = re.compile(r"\b\d{1,3}(?:\.\d+)?\s*%")


def extract_entities(text: str, known_codes=()) -> dict:
    """Regex-based extraction. Candidates only: values are not written to project data without review."""
    known = {c for c in known_codes if c in text}
    return {"project_references": sorted(known), "candidate_codes": sorted(set(_CODE.findall(text)))[:50],
            "dates": sorted(set(_DATE.findall(text)))[:50], "costs": sorted(set(m.strip() for m in _COST.findall(text)))[:50],
            "percentages": sorted(set(_PCT.findall(text)))[:50]}


def chunk(text: str, size: int = 900, overlap: int = 150) -> list:
    text = re.sub(r"[ \t]+", " ", text).strip()
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        i += size - overlap
    return out


class TfidfIndex:
    def __init__(self):
        self.chunks: list = []   # (source, text)
        self._vec = None
        self._mat = None

    def add(self, source: str, text: str):
        self.chunks += [(source, c) for c in chunk(text)]
        self._vec = None

    def _fit(self):
        from sklearn.feature_extraction.text import TfidfVectorizer
        self._vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
        self._mat = self._vec.fit_transform([c[1] for c in self.chunks])

    def search(self, query: str, k: int = 3, min_score: float = 0.08) -> list:
        if not self.chunks:
            return []
        if self._vec is None:
            self._fit()
        from sklearn.metrics.pairwise import linear_kernel
        sims = linear_kernel(self._vec.transform([query]), self._mat)[0]
        order = sims.argsort()[::-1][:k]
        return [{"source": self.chunks[i][0], "text": self.chunks[i][1], "score": round(float(sims[i]), 3)}
                for i in order if sims[i] >= min_score]
