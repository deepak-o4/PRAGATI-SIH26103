"""Ingestion, documents, reports and Copilot endpoints."""
import os
import uuid
from datetime import date
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select

from app.api.deps import ANALYSIS_ROLES, WRITE_ROLES, CurrentUser, SessionDep, audit, require
from app.copilot.documents import ALLOWED_EXT, TfidfIndex, extract_entities, extract_text
from app.copilot.engine import answer
from app.core.config import settings
from app.domain.types import DataOrigin
from app.ingestion import loader
from app.models import entities as E
from app.reports import generate as R
from app.services.analysis import recompute_all, risk_config
from app.services.repository import get_orm_project, load_orm_projects, load_projects, replace_children, to_domain, upsert_projects

router = APIRouter(tags=["data"])
MAX = settings.MAX_UPLOAD_MB * 1024 * 1024


async def _read_limited(f: UploadFile) -> bytes:
    data = await f.read(MAX + 1)
    if len(data) > MAX:
        raise HTTPException(413, f"file exceeds {settings.MAX_UPLOAD_MB} MB")
    return data


def _table_bytes(name: str, data: bytes):
    ext = Path(name or "").suffix.lower()
    if ext == ".csv" or ext == ".txt":
        return data
    if ext == ".xlsx":
        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(data), read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(it, [])]
        return [dict(zip(header, r)) for r in it if any(v is not None for v in r)]
    raise HTTPException(415, "Tabular ingestion accepts .csv or .xlsx")


@router.post("/ingestion/{kind}")
async def ingest(kind: str, db: SessionDep, file: UploadFile = File(...), origin: DataOrigin = Form(DataOrigin.USER_UPLOADED),
                 source_name: str = Form("Uploaded file"), source_url: Optional[str] = Form(None),
                 source_document: Optional[str] = Form(None), source_date: Optional[date] = Form(None),
                 snapshot_month: Optional[date] = Form(None), user=Depends(require(*WRITE_ROLES))):
    if kind not in {"projects", "snapshots", "milestones", "issues"}:
        raise HTTPException(404, "kind must be projects | snapshots | milestones | issues")
    if origin == DataOrigin.OFFICIAL and user.role not in (E.Role.SUPER_ADMIN, E.Role.ADMIN):
        raise HTTPException(403, "Only administrators may label data as OFFICIAL")
    if origin in (DataOrigin.DERIVED, DataOrigin.PREDICTED):
        raise HTTPException(422, "DERIVED/PREDICTED are system-assigned origins")
    src = _table_bytes(file.filename, await _read_limited(file))
    known = {c for c in (await db.execute(select(E.Project.project_code))).scalars().all()}
    if kind == "projects":
        projects, rep = loader.ingest_projects(src, origin, existing_codes=None)
        result_payload = projects
    elif kind == "snapshots":
        by, rep = loader.ingest_snapshots(src, known, origin)
        result_payload = by
    elif kind == "milestones":
        by, rep = loader.ingest_milestones(src, known)
        result_payload = by
    else:
        by, rep = loader.ingest_issues(src, known)
        result_payload = by
    ds = E.DataSource(source_name=source_name, source_url=source_url, source_document=source_document or file.filename,
                      source_date=source_date, snapshot_month=snapshot_month, origin=origin, uploaded_by=user.id,
                      record_counts={"total": rep.total_rows, "accepted": rep.accepted, "rejected": len(rep.rejected)},
                      report=rep.to_dict())
    db.add(ds)
    await db.flush()
    if kind == "projects":
        await upsert_projects(db, result_payload, ds.id)
    elif kind == "snapshots":
        from app.services.repository import upsert_snapshots
        for code, snaps in result_payload.items():
            await upsert_snapshots(db, code, snaps)
    else:
        for code, items in result_payload.items():
            await replace_children(db, code, **{kind: items})
    await db.flush()
    await recompute_all(db, await load_orm_projects(db))
    await audit(db, user, "INGEST", kind, ds.id, ds.record_counts)
    await db.commit()
    return {"data_source_id": str(ds.id), "report": rep.to_dict()}


@router.get("/ingestion/history")
async def ingestion_history(db: SessionDep, user: CurrentUser):
    rows = (await db.execute(select(E.DataSource).order_by(E.DataSource.created_at.desc()).limit(100))).scalars().all()
    return [{"id": str(d.id), "source_name": d.source_name, "source_url": d.source_url, "source_document": d.source_document,
             "source_date": d.source_date, "snapshot_month": d.snapshot_month, "origin": d.origin.value,
             "record_counts": d.record_counts, "created_at": d.created_at} for d in rows]


# ---- documents
def _doc_index(docs) -> TfidfIndex:
    idx = TfidfIndex()
    for d in docs:
        if d.text_content:
            idx.add(d.filename, d.text_content)
    return idx


@router.post("/documents", status_code=201)
async def upload_document(db: SessionDep, file: UploadFile = File(...), project_code: Optional[str] = Form(None),
                          user=Depends(require(*ANALYSIS_ROLES))):
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED_EXT:
        raise HTTPException(415, f"allowed types: {sorted(ALLOWED_EXT)}")
    data = await _read_limited(file)
    try:
        text = extract_text(file.filename, data)
    except Exception:
        raise HTTPException(422, "Could not extract text from this document")
    pid = None
    if project_code:
        row = await get_orm_project(db, project_code)
        if row is None:
            raise HTTPException(404, "Project not found")
        pid = row.id
    known = (await db.execute(select(E.Project.project_code))).scalars().all()
    stored = f"{uuid.uuid4().hex}{ext}"           # never trust the client filename for the path
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    (Path(settings.UPLOAD_DIR) / stored).write_bytes(data)
    d = E.ProjectDocument(project_id=pid, filename=os.path.basename(file.filename), stored_name=stored, content_type=file.content_type,
                          size_bytes=len(data), text_content=text[:2_000_000], entities=extract_entities(text, known), uploaded_by=user.id)
    db.add(d)
    await audit(db, user, "UPLOAD_DOCUMENT", "document", d.id, {"filename": d.filename})
    await db.commit()
    return {"id": str(d.id), "filename": d.filename, "entities": d.entities}


@router.get("/documents")
async def list_documents(db: SessionDep, user: CurrentUser):
    rows = (await db.execute(select(E.ProjectDocument).order_by(E.ProjectDocument.created_at.desc()).limit(200))).scalars().all()
    return [{"id": str(d.id), "filename": d.filename, "size_bytes": d.size_bytes, "entities": d.entities, "created_at": d.created_at} for d in rows]


# ---- copilot
class CopilotIn(BaseModel):
    question: str = Field(min_length=1, max_length=1000)


@router.post("/copilot")
async def copilot(body: CopilotIn, db: SessionDep, user: CurrentUser):
    ps = await load_projects(db)
    docs = (await db.execute(select(E.ProjectDocument))).scalars().all()
    idx = _doc_index(docs)
    res = answer(body.question, ps, date.today(), risk_config(), retriever=idx.search if idx.chunks else None)
    await audit(db, user, "COPILOT_QUERY", "copilot", None, {"intent": res["intent"]})
    await db.commit()
    return res


# ---- reports
@router.get("/reports/executive")
async def executive_report(db: SessionDep, user: CurrentUser, format: str = "pdf"):
    ps = await load_projects(db)
    today, cfg = date.today(), risk_config()
    acts = (await db.execute(select(E.ActionItem, E.Project.project_code).join(E.Project, E.Project.id == E.ActionItem.project_id)
                             .where(E.ActionItem.status.in_([E.ActionStatus.OPEN, E.ActionStatus.IN_PROGRESS])).limit(40))).all()
    actions = [{"title": a.title, "project_code": c, "owner": a.owner, "due_date": str(a.due_date or ""), "status": a.status.value} for a, c in acts]
    if format == "pdf":
        body, mt, ext = R.executive_pdf(ps, today, cfg, actions), "application/pdf", "pdf"
    elif format == "csv":
        body, mt, ext = R.executive_csv(ps, today, cfg), "text/csv", "csv"
    elif format == "xlsx":
        body, mt, ext = R.executive_xlsx(ps, today, cfg), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        raise HTTPException(422, "format must be pdf | csv | xlsx")
    await audit(db, user, "GENERATE_REPORT", "executive", None, {"format": format})
    await db.commit()
    return Response(body, media_type=mt, headers={"Content-Disposition": f'attachment; filename="pragati-executive-{today}.{ext}"'})


@router.get("/reports/project/{ident}")
async def project_report(ident: str, db: SessionDep, user: CurrentUser):
    row = await get_orm_project(db, ident)
    if row is None:
        raise HTTPException(404, "Project not found")
    acts = (await db.execute(select(E.ActionItem).where(E.ActionItem.project_id == row.id))).scalars().all()
    actions = [{"title": a.title, "owner": a.owner, "due_date": str(a.due_date or ""), "status": a.status.value} for a in acts]
    pdf = R.project_pdf(to_domain(row), date.today(), risk_config(), actions)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="pragati-{row.project_code}.pdf"'})
