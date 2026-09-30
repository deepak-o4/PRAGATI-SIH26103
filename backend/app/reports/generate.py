"""Executive / project review reports: PDF (reportlab), CSV, XLSX (openpyxl)."""
from __future__ import annotations

import csv
import io
from datetime import date
from typing import Optional

from app.analytics.portfolio import bottlenecks, dashboard, enrich
from app.risk.engine import compute_risk
from app.forecasting.forecast import forecast_completion

DEMO_BANNER = "Contains SYNTHETIC_DEMO data - not official government data."


def _demo(rows):
    return any(r["origin"] == "SYNTHETIC_DEMO" for r in rows)


def portfolio_rows(projects, as_of, cfg=None):
    return [enrich(p, as_of, cfg) for p in projects]


def executive_csv(projects, as_of: date, cfg=None) -> bytes:
    rows = portfolio_rows(projects, as_of, cfg)
    buf = io.StringIO()
    if not rows:
        return b""
    w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode("utf-8")


def executive_xlsx(projects, as_of: date, cfg=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    projects = list(projects)
    rows = portfolio_rows(projects, as_of, cfg)
    d = dashboard(projects, as_of, cfg)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["PRAGATI Executive Review", as_of.isoformat()])
    if _demo(rows):
        ws.append([DEMO_BANNER])
    for k, v in d["kpis"].items():
        ws.append([k, v])
    ws2 = wb.create_sheet("Projects")
    if rows:
        ws2.append(list(rows[0].keys()))
        for r in rows:
            ws2.append(list(r.values()))
        for c in ws2[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        ws2.freeze_panes = "A2"
    ws3 = wb.create_sheet("Bottlenecks")
    b = bottlenecks(projects, as_of, cfg)
    if b:
        ws3.append(list(b[0].keys()))
        for r in b:
            ws3.append(list(r.values()))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _styles():
    from reportlab.lib.styles import getSampleStyleSheet
    return getSampleStyleSheet()


def _tbl(data, widths=None):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")])]))
    return t


def _f(v, nd=1):
    return "n/a" if v is None else (f"{v:,.{nd}f}" if isinstance(v, float) else str(v))


def executive_pdf(projects, as_of: date, cfg=None, recommended_actions: Optional[list] = None) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    projects = list(projects)
    d = dashboard(projects, as_of, cfg)
    rows = portfolio_rows(projects, as_of, cfg)
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title="PRAGATI Executive Review", leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    el = [Paragraph("PRAGATI - Executive Review Report", st["Title"]), Paragraph(f"As of {as_of.isoformat()}", st["Normal"])]
    if _demo(rows):
        el.append(Paragraph(f"<b>{DEMO_BANNER}</b>", st["Normal"]))
    el += [Spacer(1, 8), Paragraph("Portfolio summary", st["Heading2"])]
    k = d["kpis"]
    el.append(_tbl([["Metric", "Value"]] + [[a.replace("_", " ").title(), _f(b, 2)] for a, b in k.items()], [260, 140]))
    el += [Paragraph("Risk distribution (active projects)", st["Heading2"]),
           _tbl([["Band", "Projects"]] + [[x["bucket"], x["count"]] for x in d["risk_distribution"]], [160, 100])]
    crit = sorted([r for r in rows if r["risk_band"] in ("CRITICAL",) and r["status"] not in ("COMPLETED", "CANCELLED")], key=lambda r: -r["risk_score"])[:15]
    el.append(Paragraph("Critical projects (top 15)", st["Heading2"]))
    el.append(_tbl([["Code", "Name", "Sector", "State", "Risk", "Progress %", "Delay d", "Overrun %"]] +
                   [[r["project_code"], r["project_name"][:40], r["sector"], r["state"], _f(r["risk_score"]), _f(r["physical_progress_pct"]),
                     _f(r["schedule_delay_days"]), _f(r["cost_overrun_pct"])] for r in crit] if crit else [["No critical projects"]]))
    ov = sorted([r for r in rows if r["cost_overrun_pct"] is not None], key=lambda r: -r["cost_overrun_pct"])[:10]
    el.append(Paragraph("Largest cost overruns", st["Heading2"]))
    el.append(_tbl([["Code", "Original Cr", "Revised Cr", "Overrun %"]] + [[r["project_code"], _f(r["original_cost_crore"]), _f(r["revised_cost_crore"], 1), _f(r["cost_overrun_pct"])] for r in ov]))
    dl = sorted([r for r in rows if (r["schedule_delay_days"] or 0) > 0], key=lambda r: -r["schedule_delay_days"])[:10]
    el.append(Paragraph("Largest schedule delays", st["Heading2"]))
    el.append(_tbl([["Code", "Total delay d", "Official slip d", "Overdue d"]] + [[r["project_code"], r["schedule_delay_days"], r["revised_slip_days"], r["overdue_days"]] for r in dl]))
    b = bottlenecks(projects, as_of, cfg)[:8]
    el.append(Paragraph("Major bottlenecks (from recorded open issues)", st["Heading2"]))
    el.append(_tbl([["Category", "Projects", "Open issues", "Avg delay d", "Avg risk", "Value Cr"]] +
                   [[x["category"], x["projects_affected"], x["open_issues"], _f(x["avg_delay_days"]), _f(x["avg_risk"]), _f(x["affected_project_value_crore"], 0)] for x in b] if b else [["No open issues recorded"]]))
    el.append(Paragraph("Recommended actions", st["Heading2"]))
    acts = recommended_actions or []
    el.append(_tbl([["Action", "Project", "Owner", "Due", "Status"]] + [[a.get("title"), a.get("project_code"), a.get("owner"), a.get("due_date"), a.get("status")] for a in acts]) if acts
               else Paragraph("No actions recorded. Actions are created by users in PRAGATI; none are auto-generated.", st["Normal"]))
    el += [Spacer(1, 8), Paragraph("Risk scores use configurable engineering weights (not validated government weights). Forecasts and scenarios are model outputs, not official projections.", st["Italic"])]
    doc.build(el)
    return buf.getvalue()


def project_pdf(p, as_of: date, cfg=None, actions: Optional[list] = None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    r = enrich(p, as_of, cfg)
    risk = compute_risk(p, as_of, cfg)
    fc = forecast_completion(p, as_of)
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"PRAGATI Project Review {p.project_code}", leftMargin=32, rightMargin=32, topMargin=32, bottomMargin=32)
    el = [Paragraph(f"Project Review - {p.project_code}", st["Title"]), Paragraph(p.project_name, st["Heading3"]),
          Paragraph(f"As of {as_of.isoformat()} | Data origin: {p.origin.value}", st["Normal"])]
    if p.origin.value == "SYNTHETIC_DEMO":
        el.append(Paragraph(f"<b>{DEMO_BANNER}</b>", st["Normal"]))
    el += [Paragraph("Project information", st["Heading2"]),
           _tbl([["Field", "Value"], ["Sector", p.sector], ["Ministry", p.line_ministry], ["Agency", p.implementing_agency], ["State", p.state], ["Status", p.status.value]], [140, 340]),
           Paragraph("Financial summary", st["Heading2"]),
           _tbl([["Item", "Crore"], ["Original cost", _f(r["original_cost_crore"])], ["Revised cost", _f(r["revised_cost_crore"])], ["Expenditure", _f(r["cumulative_expenditure_crore"])],
                 ["Cost overrun %", _f(r["cost_overrun_pct"])]], [140, 340]),
           Paragraph("Physical progress & schedule", st["Heading2"]),
           _tbl([["Item", "Value"], ["Physical progress %", _f(r["physical_progress_pct"])], ["Original end", r["original_end_date"] or "n/a"], ["Revised end (official)", r["revised_end_date"] or "n/a"],
                 ["Total slippage vs original (d)", _f(r["schedule_delay_days"])], ["Days remaining", _f(r["days_remaining"])]], [220, 260]),
           Paragraph(f"Risk: {risk.score} ({risk.band.value}) - {risk.version}", st["Heading2"]),
           _tbl([["Contributor", "Weight %", "Points", "Why"]] + [[c.label, _f(c.weight_pct), _f(c.points), c.explanation[:90]] for c in risk.contributors], [100, 50, 45, 285])]
    el.append(Paragraph("Milestones", st["Heading2"]))
    el.append(_tbl([["Milestone", "Planned end", "Actual end", "Done %", "Delay d"]] + [[m.name, m.planned_end, m.actual_end, _f(m.completion_pct, 0), m.delay_days(as_of)] for m in p.milestones] if p.milestones else [["No milestones recorded"]]))
    el.append(Paragraph("Issues", st["Heading2"]))
    el.append(_tbl([["Type", "Severity", "Status", "Description"]] + [[i.issue_type.value, i.severity.value, i.status.value, (i.description or "")[:70]] for i in p.issues] if p.issues else [["No issues recorded"]]))
    el.append(Paragraph("Forecast (PRAGATI estimate - not official)", st["Heading2"]))
    el.append(_tbl([["Item", "Value"], ["Official revised end", fc.official_revised_end_date], ["PRAGATI predicted completion", fc.predicted_completion_date or "n/a"],
                    ["Estimated delay vs official (d)", _f(fc.estimated_delay_days)], ["Confidence", f"{fc.confidence_label} ({_f(fc.confidence, 2)})"], ["Status", fc.status]], [220, 260]))
    el.append(Paragraph("Actions", st["Heading2"]))
    el.append(_tbl([["Action", "Owner", "Due", "Status"]] + [[a.get("title"), a.get("owner"), a.get("due_date"), a.get("status")] for a in actions]) if actions else Paragraph("No actions recorded.", st["Normal"]))
    doc.build(el)
    return buf.getvalue()
